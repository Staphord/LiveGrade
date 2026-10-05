"""Results export (Section 15 of the spec).

Importing students, groups and the rubric lives in ``group_import`` and
``rubric_import``; this module only writes the results workbook.
"""

from io import BytesIO

import openpyxl

PARTICIPATION_STATUS_LABELS = {
    'full': 'Fully graded', 'partial': 'Partially graded', 'none': 'Did not participate',
}


def export_results(assessment_session, results, participation_rows):
    """Returns an in-memory .xlsx workbook of participation + results.

    `participation_rows` is the same unpaginated per-student row list the
    Participation page itself renders (`_participation_rows`), so this
    sheet always matches what a lecturer sees there instead of carrying
    its own, weaker notion of "participation" — the four-status detail
    (status/graded/eligible/progress) used to live only in a since-removed
    Participation-tab export; it's centralized here now."""
    workbook = openpyxl.Workbook()

    participation_sheet = workbook.active
    participation_sheet.title = 'Participation'
    participation_sheet.append([
        'Student', 'Student ID', 'Group', 'Joined', 'Participation Status',
        'Groups Graded', 'Groups Eligible', 'Grading Progress %',
    ])
    for row in participation_rows:
        student = row['student']
        participation_sheet.append([
            student.full_name, student.student_id,
            row['own_group'].name if row['own_group'] else '',
            'Yes' if row['joined'] else 'No',
            PARTICIPATION_STATUS_LABELS.get(row['status'], row['status']),
            row['graded_count'], row['eligible_count'], row['progress_pct'],
        ])

    results_sheet = workbook.create_sheet('Results')
    results_sheet.append(['Student', 'Group', 'Group %', 'Individual %', 'Penalty %', 'Final %'])
    for row in results['student_rows']:
        results_sheet.append([
            row['student'].full_name, row['group'].name,
            float(row['group_percent']) if row['group_percent'] is not None else '',
            float(row['individual_percent']) if row['individual_percent'] is not None else '',
            row.get('penalty', 0),
            float(row['final_percent']) if row['final_percent'] is not None else '',
        ])

    group_sheet = workbook.create_sheet('Group scores')
    group_sheet.append(['Group', 'Group %'])
    for row in results['group_rows']:
        group_sheet.append([row['group'].name,
                             float(row['percent']) if row['percent'] is not None else ''])

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
