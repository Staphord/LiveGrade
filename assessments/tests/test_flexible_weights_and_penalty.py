"""One pool of 100 for the rubric's weights, no rating, basics that stay editable, and a
penalty for groups not graded that the lecturer sets."""
from decimal import Decimal
from io import BytesIO

import openpyxl
from django.core.cache import cache
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import TransactionTestCase
from django.urls import reverse

from assessments import scoring, services
from assessments.forms import RubricCategoryForm
from assessments.weight_pie import weight_pie
from assessments.models import (AssessmentSession, Evaluation, EvaluationScore, PresentationGroup,
                                PresentationTurn, RubricCategory)
from assessments.tests.test_edges_lecturer_views import LecturerTestCase, _messages


class VotedSession(LecturerTestCase):
    """Red presented and Cat graded it; Dan graded nothing, so each owes one group."""

    def setUp(self):
        super().setUp()
        cache.clear()
        self.turn = self.active_turn(self.red)
        services.submit_evaluation(
            self.turn, self.cat, {self.quality.pk: Decimal('30')},
            {self.ann.pk: {self.teamwork.pk: Decimal('20')}, self.bob.pk: {self.teamwork.pk: Decimal('40')}})
        self.set_status('live')

    def rows(self):
        return {row['student'].pk: row for row in scoring.session_results(self.session)['student_rows']}


class SessionBasicsFormTests(LecturerTestCase):

    def test_the_create_page_has_a_penalty_and_no_weight_split(self):
        page = self.client.get(reverse('assessment_session_create'))
        self.assertContains(page, 'name="ungraded_penalty"')
        self.assertContains(page, 'Penalty for each group not graded')
        self.assertNotContains(page, 'weight_percent')
        self.assertNotContains(page, '% group')

    def test_a_new_session_defaults_to_a_one_point_penalty(self):
        self.client.post(reverse('assessment_session_create'), {
            'name': 'Fresh', 'identify_by': 'full_name', 'ungraded_penalty': '1'})
        self.assertEqual(AssessmentSession.objects.get(name='Fresh').ungraded_penalty, Decimal('1'))
        self.assertEqual(AssessmentSession().ungraded_penalty, Decimal('1.00'))

    def test_a_draft_edit_changes_every_basic(self):
        self.client.post(self.url('assessment_session_edit'), {
            'name': 'New name', 'identify_by': 'full_name', 'ungraded_penalty': '3'})
        self.session.refresh_from_db()
        self.assertEqual((self.session.name, self.session.identify_by, self.session.ungraded_penalty),
                         ('New name', 'full_name', Decimal('3')))

    def test_the_basics_step_is_an_editable_form_once_the_session_exists(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal('2.5'))
        for name in ('assessment_groups', 'assessment_session_detail', 'assessment_setup'):
            page = self.client.get(self.url(name))
            self.assertContains(page, f'action="{self.url("assessment_session_edit")}"', msg_prefix=name)
            self.assertContains(page, 'name="next" value="setup"', msg_prefix=name)
            self.assertContains(page, 'id="id_basics_name"', msg_prefix=name)
            self.assertContains(page, 'id="id_basics_identify_by"', msg_prefix=name)
            self.assertContains(page, 'id="id_basics_ungraded_penalty"', msg_prefix=name)
            self.assertContains(page, 'value="2.50"', msg_prefix=name)
            self.assertContains(page, 'Save basics', msg_prefix=name)

    def test_saving_from_the_setup_step_returns_to_the_basics_step_in_draft_and_live(self):
        for status in ('draft', 'live'):
            self.set_status(status)
            response = self.client.post(self.url('assessment_session_edit'), {
                'name': f'Edited {status}', 'identify_by': 'full_name', 'ungraded_penalty': '4', 'next': 'setup'})
            self.assertRedirects(response, self.url('assessment_setup') + '?step=basics',
                                 fetch_redirect_response=False)
            self.session.refresh_from_db()
            self.assertEqual((self.session.name, self.session.ungraded_penalty),
                             (f'Edited {status}', Decimal('4')))
            self.assertEqual(self.client.get(self.url('assessment_setup') + '?step=basics').status_code, 200)

    def test_an_unknown_next_value_falls_back_to_the_normal_redirect(self):
        response = self.client.post(self.url('assessment_session_edit'), {
            'name': 'X', 'identify_by': 'full_name', 'ungraded_penalty': '1', 'next': 'http://evil.example'})
        self.assertRedirects(response, self.url('assessment_session_detail'), fetch_redirect_response=False)

    def test_an_invalid_inline_edit_shows_the_error_and_changes_nothing(self):
        response = self.client.post(self.url('assessment_session_edit'), {
            'name': 'X', 'identify_by': 'full_name', 'ungraded_penalty': '500', 'next': 'setup'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Ensure this value is less than or equal to 100')
        self.session.refresh_from_db()
        self.assertEqual(self.session.name, 'Demo day')

    def test_a_live_session_setup_page_has_the_same_editable_basics(self):
        self.set_status('live')
        page = self.client.get(self.url('assessment_setup') + '?step=basics')
        self.assertContains(page, 'id="id_basics_ungraded_penalty"')
        self.assertContains(page, 'recalculates everyone')

    def test_a_closed_session_offers_no_edit(self):
        self.set_status('closed')
        response = self.client.post(self.url('assessment_session_edit'), {
            'name': 'X', 'identify_by': 'full_name', 'ungraded_penalty': '1', 'next': 'setup'})
        self.assertRedirects(response, self.url('assessment_session_detail'), fetch_redirect_response=False)
        self.session.refresh_from_db()
        self.assertEqual(self.session.name, 'Demo day')


class FlexibleWeightPageTests(LecturerTestCase):

    def setUp(self):
        super().setUp()
        RubricCategory.objects.all().delete()

    def category(self, name, scope, weight):
        return RubricCategory.objects.create(
            assessment_session=self.session, name=name, scope=scope, weight=Decimal(weight))

    def rubric_page(self):
        return self.client.get(self.url('assessment_rubric'))

    def test_an_87_13_split_is_balanced_and_can_go_live(self):
        self.category('Project', 'group', 87)
        self.category('Teamwork', 'individual', 13)
        page = self.rubric_page()
        self.assertEqual(page.context['weight_total'], Decimal('100'))
        self.assertContains(page, 'Total weight')
        self.assertContains(page, 'Balanced')
        self.assertNotContains(page, 'class="rubric-status-banner')
        self.assertTrue(self.session.can_go_live())
        response = self.client.post(self.url('assessment_go_live'))
        self.session.refresh_from_db()
        self.assertEqual(self.session.status, 'live')
        self.assertTrue(any('now live' in m for m in _messages(response)))

    def test_a_group_only_rubric_of_100_can_go_live(self):
        self.category('Project', 'group', 100)
        self.assertTrue(self.session.can_go_live())

    def test_99_blocks_going_live_and_says_how_much_is_missing(self):
        self.category('Project', 'group', 80)
        self.category('Teamwork', 'individual', 19)
        page = self.rubric_page()
        self.assertContains(page, 'class="rubric-status-banner')
        self.assertContains(page, 'add 1% more')
        self.assertFalse(self.session.can_go_live())
        response = self.client.post(self.url('assessment_go_live'))
        self.session.refresh_from_db()
        self.assertEqual(self.session.status, 'draft')
        self.assertTrue(any('exactly 100%' in m for m in _messages(response)))

    def test_over_100_is_flagged_for_removal(self):
        self.category('Project', 'group', 80)
        self.category('Teamwork', 'individual', 30)  # saved around the form, e.g. an old draft
        page = self.rubric_page()
        self.assertContains(page, 'remove 10% of weight')
        self.assertFalse(self.session.can_go_live())

    def test_the_setup_pages_no_longer_ask_for_max_points(self):
        page = self.rubric_page()
        self.assertNotContains(page, 'Max points')
        self.assertNotContains(page, 'max_points')
        self.assertNotContains(self.client.get(self.url('assessment_rubric_table')), 'Max points')

    def test_the_add_form_shows_how_much_weight_is_free(self):
        self.category('Project', 'group', 70)
        self.assertContains(self.rubric_page(), '30% of the 100% is still free')

    def test_the_table_totals_are_one_pool_with_a_split(self):
        self.category('Project', 'group', 70)
        page = self.client.get(self.url('assessment_rubric_table'))
        self.assertContains(page, 'id="rt-total-all"')
        self.assertContains(page, 'id="rt-total-group"')
        self.assertContains(page, 'id="rt-total-individual"')
        self.assertContains(page, 'data-target="100"')

    def test_a_live_table_save_must_total_exactly_100(self):
        self.category('Project', 'group', 87)
        self.category('Teamwork', 'individual', 13)
        self.set_status('live')
        ids = list(RubricCategory.objects.order_by('pk').values_list('pk', flat=True))

        def post(w1, w2):
            data = {'rows-total': '2',
                    'rows-0-id': ids[0], 'rows-0-name': 'Project', 'rows-0-scope': 'group', 'rows-0-weight': w1,
                    'rows-1-id': ids[1], 'rows-1-name': 'Teamwork', 'rows-1-scope': 'individual', 'rows-1-weight': w2}
            return self.client.post(self.url('assessment_rubric_table'), data)

        self.assertContains(post('87', '14'), 'over 100% by 1%')
        self.assertContains(post('87', '12'), 'short by 1%')
        self.assertEqual(post('50', '50').status_code, 302)
        self.assertEqual(self.session.rubric_weight_total(), Decimal('100'))

    def test_the_live_one_at_a_time_form_stays_locked(self):
        self.category('Project', 'group', 100)
        self.set_status('live')
        response = self.client.post(self.url('assessment_rubric'), {
            'name': 'X', 'scope': 'group', 'weight': '1'})
        self.assertTrue(any('edit the rubric as a table' in m for m in _messages(response)))


class WeightPieTests(LecturerTestCase):

    def cats(self, *specs):
        RubricCategory.objects.all().delete()
        return [RubricCategory.objects.create(assessment_session=self.session, name=n, scope=sc, weight=Decimal(w))
                for n, sc, w in specs]

    def test_the_pie_has_group_individual_and_unallocated_slices_only(self):
        pie = weight_pie(self.cats(('A', 'group', 30), ('B', 'group', 20), ('C', 'individual', 25)))
        self.assertEqual([(p['name'], p['label']) for p in pie['slices']],
                         [('Group', '50'), ('Individual', '25'), ('Not allocated yet', '25')])
        self.assertEqual((pie['free'], pie['free_label'], pie['over']), (Decimal('25'), '25', False))
        self.assertEqual(pie['gradient'], '#2563eb 0.000% 50.000%, #16a34a 50.000% 75.000%, #d4d4d8 75.000% 100.000%')

    def test_twenty_categories_still_make_a_three_slice_pie(self):
        pie = weight_pie(self.cats(*[(f'C{i}', 'group' if i % 2 else 'individual', 5) for i in range(20)]))
        self.assertEqual([(p['name'], p['label']) for p in pie['slices']], [('Group', '50'), ('Individual', '50')])
        self.assertEqual(pie['free'], 0)

    def test_a_full_rubric_has_no_unallocated_slice(self):
        pie = weight_pie(self.cats(('A', 'group', 87), ('B', 'individual', '12.5'), ('C', 'individual', '0.5')))
        self.assertEqual(pie['free'], 0)
        self.assertNotIn('#d4d4d8', pie['gradient'])
        self.assertEqual([p['label'] for p in pie['slices']], ['87', '13'])

    def test_a_group_only_rubric_has_no_individual_slice(self):
        pie = weight_pie(self.cats(('A', 'group', 60)))
        self.assertEqual([p['name'] for p in pie['slices']], ['Group', 'Not allocated yet'])

    def test_an_over_full_rubric_still_makes_a_whole_circle(self):
        pie = weight_pie(self.cats(('A', 'group', 80), ('B', 'individual', 40)))
        self.assertTrue(pie['over'])
        self.assertTrue(pie['gradient'].endswith('100.000%'))

    def test_an_empty_rubric_is_one_grey_circle(self):
        pie = weight_pie([])
        self.assertEqual(pie['free_label'], '100')
        self.assertEqual([p['name'] for p in pie['slices']], ['Not allocated yet'])
        self.assertEqual(pie['gradient'], '#d4d4d8 0.000% 100.000%')

    def test_the_rubric_page_shows_the_pie_and_a_three_line_legend(self):
        self.cats(('Project', 'group', 70), ('Teamwork', 'individual', 20))
        page = self.client.get(self.url('assessment_rubric'))
        self.assertContains(page, 'id="weight-pie-card"')
        self.assertContains(page, 'conic-gradient(#2563eb')
        self.assertContains(page, 'Group</span><span class="pc">70%')
        self.assertContains(page, 'Individual</span><span class="pc">20%')
        self.assertContains(page, 'Not allocated yet</span><span class="pc">10%')
        self.assertNotContains(page, 'Project</span><span class="pc">')

    def test_a_page_with_no_categories_shows_the_all_free_legend(self):
        RubricCategory.objects.all().delete()
        page = self.client.get(self.url('assessment_rubric'))
        self.assertContains(page, 'Not allocated yet</span><span class="pc">100%')


class RubricBulkDeleteTests(LecturerTestCase):

    def setUp(self):
        super().setUp()
        self.extra = [RubricCategory.objects.create(
            assessment_session=self.session, name=f'Extra {i}', scope='group', weight=Decimal('0.5'))
            for i in range(3)]

    def delete(self, *ids):
        return self.client.post(self.url('assessment_rubric_bulk_delete'), {'category_ids': [str(i) for i in ids]})

    def test_the_rubric_step_has_select_all_and_a_box_per_category(self):
        page = self.client.get(self.url('assessment_rubric'))
        self.assertContains(page, 'id="rubric-select-all"')
        self.assertContains(page, self.url('assessment_rubric_bulk_delete'))
        self.assertContains(page, 'class="form-check-input rubric-select-box', count=5)
        for category in RubricCategory.objects.all():
            self.assertContains(page, f'value="{category.pk}" data-name="{category.name}"')

    def test_one_several_or_all_categories_can_be_deleted(self):
        response = self.delete(self.extra[0].pk)
        self.assertRedirects(response, self.url('assessment_rubric'), fetch_redirect_response=False)
        self.assertEqual(RubricCategory.objects.count(), 4)
        self.delete(self.extra[1].pk, self.extra[2].pk)
        self.assertEqual(set(RubricCategory.objects.values_list('name', flat=True)), {'Quality', 'Teamwork'})
        response = self.delete(*RubricCategory.objects.values_list('pk', flat=True))
        self.assertFalse(RubricCategory.objects.exists())
        self.assertTrue(any('Deleted 2 rubric categories' in m for m in _messages(response)))

    def test_a_single_deletion_is_worded_in_the_singular(self):
        response = self.delete(self.extra[0].pk)
        self.assertTrue(any(m == 'Deleted 1 rubric category.' for m in _messages(response)))

    def test_nothing_selected_or_junk_ids_change_nothing(self):
        for ids in ((), ('abc',), (999999,)):
            response = self.delete(*ids)
            self.assertTrue(any('Select at least one' in m for m in _messages(response)))
        self.assertEqual(RubricCategory.objects.count(), 5)

    def test_it_only_reaches_this_sessions_categories(self):
        other = AssessmentSession.objects.create(organization_id=self.org.pk, name='Other', created_by=self.lecturer)
        foreign = RubricCategory.objects.create(assessment_session=other, name='Foreign', weight=Decimal('10'))
        self.delete(foreign.pk)
        self.assertTrue(RubricCategory.objects.filter(pk=foreign.pk).exists())

    def test_the_weights_no_longer_add_up_so_the_session_cannot_go_live_until_fixed(self):
        self.delete(self.quality.pk)
        self.assertFalse(self.session.can_go_live())

    def test_a_live_session_keeps_its_rubric_and_is_pointed_to_the_table(self):
        self.set_status('live')
        response = self.delete(self.extra[0].pk)
        self.assertTrue(any('edit the rubric as a table' in m for m in _messages(response)))
        self.assertEqual(RubricCategory.objects.count(), 5)
        self.assertNotContains(self.client.get(self.url('assessment_setup') + '?step=rubric'), 'id="rubric-select-all"')

    def test_votes_stop_any_category_being_removed(self):
        self.session.status = 'draft'
        turn = PresentationTurn.objects.create(assessment_session=self.session, group=self.red, status='closed')
        Evaluation.objects.create(presentation_turn=turn, evaluator=self.cat)
        response = self.delete(self.extra[0].pk)
        self.assertTrue(any('cannot be added or removed' in m for m in _messages(response)))
        self.assertEqual(RubricCategory.objects.count(), 5)

    def test_a_closed_session_is_read_only(self):
        self.set_status('closed')
        response = self.delete(self.extra[0].pk)
        self.assertRedirects(response, self.url('assessment_session_detail'), fetch_redirect_response=False)
        self.assertEqual(RubricCategory.objects.count(), 5)

    def test_get_is_not_allowed(self):
        self.assertEqual(self.client.get(self.url('assessment_rubric_bulk_delete')).status_code, 405)


class EmptyWeightFieldTests(LecturerTestCase):

    def test_the_new_category_form_leaves_weight_empty_with_a_placeholder(self):
        form = RubricCategoryForm(assessment_session=self.session)
        html = str(form['weight'])
        self.assertIn('placeholder="e.g. 30"', html)
        self.assertNotIn('value=', html)
        page = self.client.get(self.url('assessment_rubric'))
        self.assertNotContains(page, 'name="weight" value="0.00"')
        self.assertContains(page, 'id="id_rubric_weight"')

    def test_editing_an_existing_category_keeps_its_weight(self):
        form = RubricCategoryForm(instance=self.quality, assessment_session=self.session)
        self.assertIn('value="60"', str(form['weight']))

    def test_a_blank_weight_is_refused(self):
        form = RubricCategoryForm({'name': 'x', 'scope': 'group', 'weight': ''}, assessment_session=self.session)
        self.assertFalse(form.is_valid())
        self.assertIn('weight', form.errors)

    def test_a_new_row_in_the_table_starts_empty_with_a_placeholder(self):
        page = self.client.get(self.url('assessment_rubric_table'))
        self.assertContains(page, 'name="rows-__i__-weight" value="" placeholder="e.g. 30"')


class StudentScoreRangeTests(VotedSession):

    def test_the_evaluate_page_scores_each_category_from_zero_to_its_weight(self):
        PresentationTurn.objects.all().delete()
        Evaluation.objects.all().delete()
        self.active_turn(self.blue)
        self.client.post(reverse('student_join', args=[self.session.uuid]), {'identifier': self.ann.student_id})
        page = self.client.get(reverse('student_evaluate', args=[self.session.uuid]))
        self.assertContains(page, 'max="60.000000"')
        self.assertContains(page, 'out of 60')
        self.assertContains(page, 'out of 40')
        self.assertNotContains(page, 'max_points')

    def test_a_fractional_weight_is_shown_without_padding(self):
        RubricCategory.objects.filter(pk=self.quality.pk).update(weight=Decimal('12.5'))
        PresentationTurn.objects.all().delete()
        Evaluation.objects.all().delete()
        self.active_turn(self.blue)
        self.client.post(reverse('student_join', args=[self.session.uuid]), {'identifier': self.ann.student_id})
        page = self.client.get(reverse('student_evaluate', args=[self.session.uuid]))
        self.assertContains(page, 'out of 12.5')

    def test_a_vote_is_stored_with_the_weight_it_was_given_out_of(self):
        self.assertEqual(set(EvaluationScore.objects.values_list('max_value', flat=True)),
                         {Decimal('60'), Decimal('40')})


class PenaltyTests(VotedSession):

    def test_the_penalty_is_per_ungraded_group_at_the_sessions_rate(self):
        # Dan (blue) never graded red: one group owed.
        for rate in ('0', '1', '2.5'):
            AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal(rate))
            self.session.refresh_from_db()
            row = self.rows()[self.dan.pk]
            self.assertEqual((row['ungraded'], row['penalty']), (1, Decimal(rate)))
            expected = max(Decimal('0'), (row['group_percent'] or 0) + (row['individual_percent'] or 0) - Decimal(rate))
            self.assertEqual(row['final_percent'], expected.quantize(Decimal('0.1')))
            self.assertEqual(scoring.final_score_percent(self.dan, self.blue, self.session), row['final_percent'])

    def test_a_student_who_graded_everything_owes_nothing(self):
        row = self.rows()[self.cat.pk]
        self.assertEqual((row['ungraded'], row['penalty']), (0, 0))

    def test_changing_the_penalty_while_live_rewrites_past_results(self):
        # Ann and Bob are in Red and graded nobody else: they owe Blue (which hasn't presented).
        before = self.rows()[self.ann.pk]['final_percent']
        self.client.post(self.url('assessment_session_edit'), {
            'name': 'Demo day', 'identify_by': 'student_id', 'ungraded_penalty': '5'})
        self.session.refresh_from_db()
        after = self.rows()[self.ann.pk]['final_percent']
        self.assertEqual(before - after, Decimal('4.0'))
        self.assertEqual(self.session.changes.get().summary, 'Changed the penalty per group not graded.')

    def test_the_final_never_goes_below_zero(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal('100'))
        self.session.refresh_from_db()
        self.assertEqual(self.rows()[self.dan.pk]['final_percent'], Decimal('0.0'))
        self.assertEqual(scoring.final_score_percent(self.dan, self.blue, self.session), Decimal('0.0'))

    def test_the_results_page_names_the_sessions_penalty_not_one_percent(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal('2.5'))
        page = self.client.get(self.url('assessment_results'))
        self.assertContains(page, '2.50 per other group not graded')
        self.assertContains(page, '2.50 points deducted per other group')
        self.assertNotContains(page, '− 1% per other group')
        self.assertContains(page, '−2.50%')

    def test_the_results_page_hides_the_badge_when_there_is_no_penalty(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal('0'))
        page = self.client.get(self.url('assessment_results'))
        self.assertNotContains(page, 'badge bg-danger-subtle')

    def test_the_penalty_column_can_be_sorted(self):
        self.assertEqual(self.client.get(self.url('assessment_results') + '?sort=penalty&dir=desc').status_code, 200)

    def test_the_export_carries_the_points_deducted(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(ungraded_penalty=Decimal('2.5'))
        response = self.client.get(self.url('assessment_results_export'))
        book = openpyxl.load_workbook(BytesIO(response.content))
        sheet = book['Results']
        self.assertEqual(sheet.cell(1, 5).value, 'Penalty (points)')
        dan = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == 'Dan Dyer')
        self.assertEqual(dan[4], 2.5)


class ScoreDisplayTests(VotedSession):

    def test_group_results_are_points_of_the_group_categories(self):
        # Cat gave 30 of 60 to Red; Ann got 20/40, Bob 40/40 from Cat.
        rows = self.rows()
        self.assertEqual(rows[self.ann.pk]['group_percent'], Decimal('30.0'))
        self.assertEqual(rows[self.ann.pk]['individual_percent'], Decimal('20.0'))
        self.assertEqual(rows[self.bob.pk]['individual_percent'], Decimal('40.0'))

    def test_the_results_chart_is_capped_at_the_group_categories_weight(self):
        page = self.client.get(self.url('assessment_results'))
        self.assertContains(page, 'max: 60.00')

    def test_the_chart_falls_back_to_100_with_no_group_categories(self):
        RubricCategory.objects.filter(scope='group').delete()
        page = self.client.get(self.url('assessment_results'))
        self.assertContains(page, 'max: 100.00')

    def test_the_live_rankings_bar_scales_to_the_group_weight_and_survives_none(self):
        page = self.client.get(self.url('assessment_session_detail'))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'width:50%')  # 30 of 60
        RubricCategory.objects.filter(scope='group').delete()
        self.assertEqual(self.client.get(self.url('assessment_session_detail')).status_code, 200)


class MigrationTests(TransactionTestCase):
    """0004 keeps every past score meaning what it did: each is out of its category's old max."""

    migrate_from = [('assessments', '0003_group_topic_location_session_change')]
    migrate_to = [('assessments', '0004_flexible_weights_penalty')]

    def test_old_scores_keep_their_scale_and_old_sessions_get_the_one_point_penalty(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_from)
        old = executor.loader.project_state(self.migrate_from).apps
        Session, Category = old.get_model('assessments', 'AssessmentSession'), old.get_model('assessments', 'RubricCategory')
        Group, Student = old.get_model('assessments', 'PresentationGroup'), old.get_model('assessments', 'Student')
        Turn, Evaluation_ = old.get_model('assessments', 'PresentationTurn'), old.get_model('assessments', 'Evaluation')
        Score = old.get_model('assessments', 'EvaluationScore')
        session = Session.objects.create(organization_id=1, name='Old', group_weight_percent=60, individual_weight_percent=40)
        category = Category.objects.create(assessment_session=session, name='Q', max_points=Decimal('20'), weight=Decimal('60'))
        other = Category.objects.create(assessment_session=session, name='R', max_points=Decimal('5'), weight=Decimal('40'))
        group = Group.objects.create(assessment_session=session, name='G')
        student = Student.objects.create(assessment_session=session, full_name='S', student_id='S1')
        turn = Turn.objects.create(assessment_session=session, group=group)
        evaluation = Evaluation_.objects.create(presentation_turn=turn, evaluator=student)
        a = Score.objects.create(evaluation=evaluation, rubric_category=category, value=Decimal('15'))
        b = Score.objects.create(evaluation=evaluation, rubric_category=other, value=Decimal('4.5'))

        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_to)
        new = executor.loader.project_state(self.migrate_to).apps
        self.assertEqual(new.get_model('assessments', 'EvaluationScore').objects.get(pk=a.pk).max_value, Decimal('20'))
        self.assertEqual(new.get_model('assessments', 'EvaluationScore').objects.get(pk=b.pk).max_value, Decimal('5'))
        self.assertEqual(new.get_model('assessments', 'AssessmentSession').objects.get(pk=session.pk).ungraded_penalty,
                         Decimal('1.00'))
        fields = {f.name for f in new.get_model('assessments', 'RubricCategory')._meta.get_fields()}
        self.assertNotIn('max_points', fields)
        self.assertNotIn('group_weight_percent',
                         {f.name for f in new.get_model('assessments', 'AssessmentSession')._meta.get_fields()})
