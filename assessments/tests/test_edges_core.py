"""The assessment engine's edges: turn races, scheduled jobs, scoring with
nothing to score, the spreadsheet importers, and what a student's cookie
is and is not allowed to mean.

Students have no account. The only thing identifying one is a cookie holding
roster ids, so the tests that matter most here are the ones saying what that
cookie cannot do: act in another session, vote twice, vote for a category that
is not in the rubric, or score outside the range the rubric allows.
"""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

import openpyxl
from accounts.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, SimpleTestCase, TestCase
from django.utils import timezone

from accounts.testing import make_org
from assessments import excel, realtime, scoring, services, tasks, turns
from assessments.forms import AssessmentSessionForm, RubricCategoryForm
from assessments.models import (AssessmentSession, Evaluation, EvaluationScore,
                                GroupMembership, ParticipationRecord,
                                PresentationGroup, PresentationTurn,
                                RubricCategory, Student)
from assessments.pagination import paginate


def workbook(*rows):
    """An in-memory .xlsx holding ``rows``, first row the header."""
    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(list(row))
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


class SessionFixture(TestCase):
    """A session with two groups of two students and one rubric per scope."""

    def setUp(self):
        self.org = make_org('Uni', slug='uni')
        self.session = AssessmentSession.objects.create(
            organization_id=self.org.pk, name='Demo day',
            identify_by=AssessmentSession.IdentifyBy.STUDENT_ID)
        self.ann = self.student('Ann Able', 'S1')
        self.bob = self.student('Bob Best', 'S2')
        self.cat = self.student('Cat Cole', 'S3')
        self.dan = self.student('Dan Dyer', 'S4')
        self.red = self.group('Red', 1, self.ann, self.bob)
        self.blue = self.group('Blue', 2, self.cat, self.dan)
        self.quality = RubricCategory.objects.create(
            assessment_session=self.session, name='Quality',
            scope=RubricCategory.Scope.GROUP, weight=60)
        self.teamwork = RubricCategory.objects.create(
            assessment_session=self.session, name='Teamwork',
            scope=RubricCategory.Scope.INDIVIDUAL, weight=40)

    def student(self, name, student_id, session=None):
        return Student.objects.create(
            assessment_session=session or self.session, full_name=name, student_id=student_id)

    def group(self, name, order, *members):
        group = PresentationGroup.objects.create(
            assessment_session=self.session, name=name, order=order)
        for member in members:
            GroupMembership.objects.create(group=group, student=member)
        return group

    def active_turn(self, group, **fields):
        now = timezone.now()
        return PresentationTurn.objects.create(
            assessment_session=self.session, group=group,
            status=PresentationTurn.Status.ACTIVE, opened_at=now,
            voting_opened_at=now, **fields)


class TurnEdgeTests(SessionFixture):

    def test_activating_the_designated_next_group_clears_the_designation(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(
            designated_next_group=self.blue)
        turns.activate_turn(self.session, self.blue)
        self.session.refresh_from_db()
        self.assertIsNone(self.session.designated_next_group)

    def test_opening_voting_on_a_turn_that_is_not_active_changes_nothing(self):
        closed = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.CLOSED)
        self.assertEqual(turns.open_voting(closed), closed)
        closed.refresh_from_db()
        self.assertIsNone(closed.voting_opened_at)

    @mock.patch('assessments.tasks.auto_advance_turn.apply_async')
    def test_opening_voting_with_a_length_schedules_the_close(self, schedule):
        turn = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.ACTIVE, opened_at=timezone.now())
        turns.open_voting(turn, duration_seconds=60)
        self.assertIsNotNone(turn.voting_ends_at)
        schedule.assert_called_once()

    def test_opening_voting_with_no_length_leaves_it_open_ended(self):
        turn = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.ACTIVE, opened_at=timezone.now())
        turns.open_voting(turn)
        self.assertIsNone(turn.voting_ends_at)

    def test_a_transition_with_nothing_pending_does_nothing(self):
        self.assertIsNone(turns.activate_pending_transition(self.session))
        self.assertFalse(PresentationTurn.objects.exists())

    def test_a_transition_whose_gap_has_not_finished_waits(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(
            pending_next_group=self.blue,
            pending_transition_at=timezone.now() + timedelta(minutes=5))
        turns.activate_pending_transition(self.session)
        self.assertFalse(PresentationTurn.objects.exists())

    def test_a_transition_another_caller_already_claimed_is_not_run_twice(self):
        self.session.pending_next_group = self.blue
        self.session.pending_transition_at = timezone.now() - timedelta(seconds=1)
        with mock.patch.object(AssessmentSession, 'refresh_from_db'):
            turns.activate_pending_transition(self.session)    # DB has nothing pending
        self.assertFalse(PresentationTurn.objects.exists())

    def test_closing_a_turn_that_is_already_closed_reports_losing_the_race(self):
        turn = self.active_turn(self.red)
        self.assertTrue(turns.close_turn(turn))
        self.assertFalse(turns.close_turn(turn))

    def test_auto_advance_backs_off_when_another_caller_closed_the_turn(self):
        turn = self.active_turn(self.red, voting_ends_at=timezone.now() - timedelta(seconds=1))
        with mock.patch.object(turns, 'close_turn', return_value=False), \
                mock.patch.object(turns, '_advance_to') as advance:
            turns.auto_advance(turn)
        advance.assert_not_called()

    def test_an_adjustment_that_runs_the_timer_out_does_not_advance_twice(self):
        turn = self.active_turn(self.red, voting_ends_at=timezone.now() + timedelta(seconds=5))
        with mock.patch.object(turns, 'close_turn', return_value=False), \
                mock.patch.object(turns, '_advance_to') as advance:
            self.assertIsNone(turns.adjust_timer(turn, -60))
        advance.assert_not_called()


class ScheduledJobTests(SessionFixture):

    def test_a_scheduled_voting_open_opens_a_presenting_turn(self):
        turn = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.ACTIVE, opened_at=timezone.now())
        with mock.patch('assessments.turns.open_voting') as open_voting:
            tasks.open_scheduled_voting(turn.pk)
        open_voting.assert_called_once()

    def test_a_scheduled_voting_open_leaves_a_turn_that_already_opened(self):
        turn = self.active_turn(self.red)
        with mock.patch('assessments.turns.open_voting') as open_voting:
            tasks.open_scheduled_voting(turn.pk)
        open_voting.assert_not_called()

    def test_the_voting_window_closing_is_announced_once_it_has_passed(self):
        turn = self.active_turn(self.red, voting_ends_at=timezone.now() - timedelta(seconds=1))
        with mock.patch('assessments.realtime.broadcast_session_event') as broadcast:
            tasks.close_voting_window(turn.pk)
        self.assertEqual(broadcast.call_args.args[1], 'voting.closed')

    def test_a_voting_window_that_was_extended_is_not_announced_early(self):
        turn = self.active_turn(self.red, voting_ends_at=timezone.now() + timedelta(minutes=5))
        with mock.patch('assessments.realtime.broadcast_session_event') as broadcast:
            tasks.close_voting_window(turn.pk)
        broadcast.assert_not_called()

    def test_finishing_a_presentation_advances_the_session(self):
        turn = self.active_turn(self.red)
        with mock.patch('assessments.turns.auto_advance') as advance:
            tasks.finish_presentation(turn.pk)
        advance.assert_called_once()
        turn.refresh_from_db()
        self.assertIsNotNone(turn.voting_ends_at)

    def test_finishing_a_presentation_that_already_closed_does_nothing(self):
        turn = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.CLOSED)
        with mock.patch('assessments.turns.auto_advance') as advance:
            tasks.finish_presentation(turn.pk)
        advance.assert_not_called()

    def test_an_auto_advance_for_a_turn_that_is_gone_is_ignored(self):
        self.assertIsNone(tasks.auto_advance_turn(999999))

    def test_an_auto_advance_that_fails_is_retried(self):
        turn = self.active_turn(self.red)
        with mock.patch('assessments.turns.auto_advance', side_effect=RuntimeError('boom')):
            with self.assertRaises(RuntimeError):
                tasks.auto_advance_turn(turn.pk)

    def test_a_pending_group_for_a_session_that_is_gone_is_ignored(self):
        self.assertIsNone(tasks.activate_pending_group(999999))

    def test_a_pending_group_that_fails_to_open_is_retried(self):
        with mock.patch('assessments.turns.activate_pending_transition',
                        side_effect=RuntimeError('boom')):
            with self.assertRaises(RuntimeError):
                tasks.activate_pending_group(self.session.pk)


class ServiceTests(SessionFixture):

    def request(self, ids=None):
        request = RequestFactory().get('/')
        request.session = {services.STUDENT_SESSION_KEY: ids} if ids is not None else {}
        return request

    def test_a_blank_answer_matches_nobody(self):
        for typed in (None, '', '   '):
            self.assertIsNone(services.match_student(self.session, typed))

    def test_matching_ignores_case_and_surrounding_space(self):
        self.assertEqual(services.match_student(self.session, '  s1 '), self.ann)

    def test_a_name_session_matches_on_the_name(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(identify_by='full_name')
        self.session.refresh_from_db()
        self.assertEqual(services.match_student(self.session, 'bob best'), self.bob)

    def test_joining_twice_does_not_duplicate_the_cookie_or_the_record(self):
        request = self.request()
        services.record_join(request, self.ann)
        services.record_join(request, self.ann)
        self.assertEqual(request.session[services.STUDENT_SESSION_KEY], [self.ann.pk])
        self.assertEqual(ParticipationRecord.objects.count(), 1)

    def test_a_cookie_naming_a_student_of_another_session_means_nobody(self):
        other_session = AssessmentSession.objects.create(organization_id=self.org.pk, name='Other')
        outsider = self.student('Out Sider', 'X1', session=other_session)
        self.assertIsNone(services.student_from_request(self.request([outsider.pk]), self.session))
        self.assertEqual(
            services.student_from_request(self.request([outsider.pk]), other_session), outsider)

    def test_no_cookie_means_nobody(self):
        self.assertIsNone(services.student_from_request(self.request(), self.session))

    def test_a_score_for_a_category_that_is_not_in_the_rubric_is_refused(self):
        turn = self.active_turn(self.red)
        with self.assertRaisesMessage(ValidationError, 'Unknown rubric category'):
            services.submit_evaluation(turn, self.cat, {999999: Decimal('5')}, {})
        self.assertFalse(Evaluation.objects.exists())

    def test_a_group_score_cannot_be_filed_under_an_individual_category(self):
        turn = self.active_turn(self.red)
        with self.assertRaisesMessage(ValidationError, 'Unknown rubric category'):
            services.submit_evaluation(turn, self.cat, {self.teamwork.pk: Decimal('5')}, {})
        self.assertFalse(Evaluation.objects.exists())

    def test_a_score_outside_the_rubrics_range_is_refused_and_nothing_is_kept(self):
        turn = self.active_turn(self.red)
        for bad in (Decimal('60.01'), Decimal('-1')):
            with self.subTest(score=bad):
                with self.assertRaises(ValidationError):
                    services.submit_evaluation(turn, self.cat, {self.quality.pk: bad}, {})
        self.assertFalse(Evaluation.objects.exists())

    def test_a_member_of_the_presenting_group_cannot_grade_it(self):
        turn = self.active_turn(self.red)
        with self.assertRaisesMessage(ValidationError, 'cannot evaluate this group'):
            services.submit_evaluation(turn, self.ann, {self.quality.pk: Decimal('5')}, {})

    def test_nobody_votes_twice(self):
        turn = self.active_turn(self.red)
        services.submit_evaluation(turn, self.cat, {self.quality.pk: Decimal('5')}, {})
        with self.assertRaisesMessage(ValidationError, 'already submitted'):
            services.submit_evaluation(turn, self.cat, {self.quality.pk: Decimal('9')}, {})
        self.assertEqual(EvaluationScore.objects.get().value, Decimal('5'))

    def test_nobody_votes_while_voting_is_paused(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(voting_paused=True)
        turn = self.active_turn(self.red)
        turn.refresh_from_db()
        with self.assertRaisesMessage(ValidationError, 'paused'):
            services.submit_evaluation(turn, self.cat, {self.quality.pk: Decimal('5')}, {})


class ScoringEdgeTests(SessionFixture):

    def test_a_group_nobody_has_scored_has_no_score(self):
        self.assertIsNone(scoring.group_score_percent(self.red, self.session))
        self.assertIsNone(scoring.individual_score_percent(self.ann, self.session))

    def test_a_session_with_no_categories_has_no_score(self):
        RubricCategory.objects.all().delete()
        self.assertIsNone(scoring.group_score_percent(self.red, self.session))
        self.assertIsNone(scoring.individual_score_percent(self.ann, self.session))

    def test_a_category_nobody_scored_is_skipped_not_counted_as_zero(self):
        RubricCategory.objects.create(
            assessment_session=self.session, name='Extra',
            scope=RubricCategory.Scope.GROUP, weight=0)
        turn = self.active_turn(self.red)
        services.submit_evaluation(turn, self.cat, {self.quality.pk: Decimal('60')}, {})
        self.assertEqual(scoring.group_score_percent(self.red, self.session), Decimal('60.0'))

    def test_every_other_group_a_student_skipped_costs_a_point(self):
        self.assertEqual(scoring.ungraded_group_count(self.ann, self.red, self.session), 1)
        turn = self.active_turn(self.blue)
        services.submit_evaluation(turn, self.ann, {self.quality.pk: Decimal('5')}, {})
        self.assertEqual(scoring.ungraded_group_count(self.ann, self.red, self.session), 0)

    def test_a_session_with_one_group_has_nothing_to_skip(self):
        self.blue.delete()
        self.assertEqual(scoring.ungraded_group_count(self.ann, self.red, self.session), 0)

    def test_the_final_score_never_goes_below_zero(self):
        for n in range(1, 4):
            self.group(f'Extra {n}', 10 + n)
        # No scores at all, three more groups skipped: 0 - 4 would be negative.
        self.assertEqual(scoring.final_score_percent(self.ann, self.red, self.session),
                         Decimal('0.0'))


class ModelEdgeTests(SessionFixture):

    def test_names(self):
        self.assertEqual(str(self.session), 'Demo day')
        self.assertEqual(str(self.quality), 'Quality (Group / project, 60%)')
        self.assertEqual(str(self.red), 'Red')
        turn = self.active_turn(self.red)
        self.assertEqual(str(turn), 'Red — Active')
        record = ParticipationRecord.objects.create(
            assessment_session=self.session, student=self.ann)
        self.assertEqual(str(record), 'Ann Able joined Demo day')
        group_vote = Evaluation.objects.create(presentation_turn=turn, evaluator=self.cat)
        self.assertEqual(str(group_vote), f'Cat Cole on {turn}')
        single = Evaluation.objects.create(
            presentation_turn=turn, evaluator=self.dan, target_student=self.ann)
        self.assertEqual(str(single), f'Dan Dyer on {turn} -> Ann Able')
        score = EvaluationScore.objects.create(
            evaluation=group_vote, rubric_category=self.quality, value=7)
        self.assertEqual(str(score), 'Quality: 7')

    def test_a_name_session_needs_a_name_on_every_roster_row(self):
        AssessmentSession.objects.filter(pk=self.session.pk).update(identify_by='full_name')
        self.session.refresh_from_db()
        with self.assertRaisesMessage(ValidationError, 'identifies students by full name'):
            Student(assessment_session=self.session, full_name='', student_id='S9').clean()

    def test_a_student_with_no_session_is_not_checked(self):
        self.assertIsNone(Student(full_name='x').clean())

    def test_a_membership_with_nothing_to_check_is_not_checked(self):
        self.assertIsNone(GroupMembership().clean())

    def test_revalidating_an_existing_membership_does_not_clash_with_itself(self):
        membership = GroupMembership.objects.get(student=self.ann)
        membership.clean()

    def test_a_student_cannot_belong_to_two_groups(self):
        with self.assertRaisesMessage(ValidationError, 'already in "Red"'):
            GroupMembership(group=self.blue, student=self.ann).clean()

    def test_a_full_group_takes_no_more_members(self):
        PresentationGroup.objects.filter(pk=self.red.pk).update(max_size=2)
        self.red.refresh_from_db()
        extra = self.student('Eve Ede', 'S5')
        with self.assertRaisesMessage(ValidationError, 'already full'):
            GroupMembership(group=self.red, student=extra).clean()

    def test_voting_is_not_open_while_the_group_is_still_presenting(self):
        turn = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.ACTIVE,
            presentation_ends_at=timezone.now() + timedelta(minutes=5))
        self.assertFalse(turn.is_voting_open)

    def test_voting_is_not_open_once_its_deadline_has_passed(self):
        turn = self.active_turn(self.red, voting_ends_at=timezone.now() - timedelta(seconds=1))
        self.assertFalse(turn.is_voting_open)

    def test_voting_is_open_before_its_deadline_and_for_an_untimed_turn(self):
        self.assertTrue(self.active_turn(self.red).is_voting_open)
        self.assertTrue(self.active_turn(
            self.blue, voting_ends_at=timezone.now() + timedelta(minutes=1)).is_voting_open)

    def test_the_ungraded_penalty_must_be_between_zero_and_a_hundred(self):
        for bad in (Decimal('-0.01'), Decimal('100.01')):
            self.session.ungraded_penalty = bad
            with self.assertRaises(ValidationError):
                self.session.full_clean(exclude=['created_by'])
        for good in (Decimal('0'), Decimal('2.5'), Decimal('100')):
            self.session.ungraded_penalty = good
            self.session.full_clean(exclude=['created_by'])


class FormEdgeTests(SessionFixture):

    def test_the_penalty_form_accepts_zero_and_rejects_negative_or_over_a_hundred(self):
        base = {'name': 'x', 'identify_by': 'full_name'}
        self.assertTrue(AssessmentSessionForm({**base, 'ungraded_penalty': '0'}).is_valid())
        self.assertTrue(AssessmentSessionForm({**base, 'ungraded_penalty': '2.5'}).is_valid())
        self.assertFalse(AssessmentSessionForm({**base, 'ungraded_penalty': '-1'}).is_valid())
        self.assertFalse(AssessmentSessionForm({**base, 'ungraded_penalty': '101'}).is_valid())
        self.assertFalse(AssessmentSessionForm(base).is_valid())

    def test_a_category_form_with_no_session_has_no_budget_to_check(self):
        form = RubricCategoryForm({'name': 'x', 'scope': 'group', 'weight': '50'})
        self.assertTrue(form.is_valid(), form.errors)

    def test_a_weight_must_be_more_than_zero_and_at_most_a_hundred(self):
        for bad in ('0', '-5', '100.01', '500'):
            form = RubricCategoryForm({'name': 'x', 'scope': 'group', 'weight': bad})
            self.assertFalse(form.is_valid(), bad)
            self.assertIn('weight', form.errors)


class PaginationAndAdminTests(SimpleTestCase):

    def page(self, **params):
        request = RequestFactory().get('/', params)
        return paginate(request, list(range(60)))

    def test_a_junk_or_unoffered_size_falls_back_to_the_default(self):
        self.assertEqual(self.page(per_page='abc')[0].per_page, 25)
        self.assertEqual(self.page(per_page='7')[0].per_page, 25)
        self.assertEqual(self.page(per_page='50')[0].per_page, 50)

    def test_a_junk_page_is_the_first_and_a_page_past_the_end_is_the_last(self):
        self.assertEqual(self.page(page='x')[1].number, 1)
        self.assertEqual(self.page(page='99')[1].number, 3)

    def test_with_no_live_channel_a_broadcast_is_dropped_quietly(self):
        with mock.patch.object(realtime, 'get_channel_layer', return_value=None):
            self.assertIsNone(realtime.broadcast_session_event(mock.Mock(), 'x'))


class ExportTests(SessionFixture):

    def test_the_export_has_one_sheet_each_for_participation_results_and_groups(self):
        participation = [{
            'student': self.ann, 'own_group': self.red, 'joined': True, 'status': 'full',
            'graded_count': 1, 'eligible_count': 1, 'progress_pct': 100,
        }, {
            'student': self.bob, 'own_group': None, 'joined': False, 'status': 'mystery',
            'graded_count': 0, 'eligible_count': 1, 'progress_pct': 0,
        }]
        results = {
            'student_rows': [
                {'student': self.ann, 'group': self.red, 'group_percent': Decimal('50.0'),
                 'individual_percent': None, 'final_percent': Decimal('50.0')},
                {'student': self.bob, 'group': self.red, 'group_percent': None,
                 'individual_percent': Decimal('10.0'), 'final_percent': None, 'penalty': Decimal('2.5')},
            ],
            'group_rows': [{'group': self.red, 'percent': Decimal('50.0')},
                           {'group': self.blue, 'percent': None}],
        }
        book = openpyxl.load_workbook(excel.export_results(self.session, results, participation))
        self.assertEqual(book.sheetnames, ['Participation', 'Results', 'Group scores'])
        sheet = book['Participation']
        self.assertEqual(sheet.cell(2, 5).value, 'Fully graded')
        self.assertEqual(sheet.cell(3, 5).value, 'mystery')
        self.assertEqual(sheet.cell(3, 3).value, None)
        self.assertEqual(book['Results'].cell(2, 4).value, None)
        self.assertEqual(book['Group scores'].cell(3, 2).value, None)


class VotingWindowNoOpTests(SessionFixture):

    def test_a_voting_window_that_has_nothing_to_close_is_left_alone(self):
        closed = PresentationTurn.objects.create(
            assessment_session=self.session, group=self.red,
            status=PresentationTurn.Status.CLOSED)
        untimed = self.active_turn(self.blue)
        with mock.patch('assessments.realtime.broadcast_session_event') as broadcast:
            tasks.close_voting_window(closed.pk)
            tasks.close_voting_window(untimed.pk)
        broadcast.assert_not_called()


class UploadSafetyTests(SessionFixture):
    """Spreadsheets arrive from a browser. Whatever is uploaded, the importer
    answers with a message; it never falls over, and it never even opens a file
    that is the wrong kind or an unreasonable size."""

    def test_a_file_that_is_not_a_workbook_is_reported_not_raised(self):
        from assessments import group_import, rubric_import
        from assessments.sheets import UNREADABLE

        junk = BytesIO(b'this is not a workbook')
        self.assertEqual(group_import.parse_groups(junk)['fatal'], UNREADABLE)
        self.assertEqual(group_import.parse_groups(BytesIO(b'PK\x03\x04 truncated'))['fatal'], UNREADABLE)
        self.assertEqual(rubric_import.parse_rubric(BytesIO(b'nope'))['fatal'], UNREADABLE)

    def test_only_a_modest_xlsx_is_accepted(self):
        from assessments.forms import MAX_IMPORT_BYTES, GroupImportForm, RubricImportForm

        def accepted(form_class, name, size=10):
            upload = SimpleUploadedFile(name, b'x' * size)
            data = {'mode': 'replace'} if form_class is RubricImportForm else {}
            return form_class(data, {'file': upload}).is_valid()

        for form_class in (GroupImportForm, RubricImportForm):
            self.assertTrue(accepted(form_class, 'groups.xlsx'))
            self.assertTrue(accepted(form_class, 'GROUPS.XLSX'))
            self.assertFalse(accepted(form_class, 'groups.xls'))
            self.assertFalse(accepted(form_class, 'groups.exe'))
            self.assertFalse(accepted(form_class, 'groups'))
            self.assertFalse(accepted(form_class, 'big.xlsx', size=MAX_IMPORT_BYTES + 1))
            self.assertTrue(accepted(form_class, 'edge.xlsx', size=MAX_IMPORT_BYTES))
