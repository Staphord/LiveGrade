"""Shared pieces for reading the spreadsheets lecturers upload."""
import re
from io import BytesIO

import openpyxl

UNREADABLE = 'That file could not be read as an Excel workbook.'

#: Names and IDs share one cell, separated by a comma, a semicolon or a line break.
SEPARATOR = re.compile(r'[,;\r\n]+')


def clean(value):
    """A cell as one tidy line: no leading/trailing space, runs of spaces collapsed."""
    return ' '.join(cell_text(value).split())


def cell_text(value):
    """A cell as text. A number typed into Excel arrives as 1001.0; show it as 1001."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def key(value):
    """What makes two spellings of a name the same person."""
    return clean(value).casefold()


def split_people(value):
    """``'Ann Lee, Bo Ng,Cy Oh'`` -> ``['Ann Lee', 'Bo Ng', 'Cy Oh']``. Empty pieces vanish."""
    return [piece for piece in (clean(part) for part in SEPARATOR.split(cell_text(value))) if piece]


def _value(cell, percent_as_points):
    value = cell.value
    if (percent_as_points and isinstance(value, (int, float)) and not isinstance(value, bool)
            and str(getattr(cell, 'number_format', '')).endswith('%')):
        return round(value * 100, 6)
    return value


def read_rows(uploaded_file, percent_as_points=False):
    """The first sheet's rows as lists of values, or ``None`` if it is not a workbook.

    ``percent_as_points``: a cell formatted as a percentage holds 0.3 for "30%";
    read it as 30, which is what the person typed and means.

    A corrupt or mislabelled upload raises a different exception from each layer
    (zip, XML, openpyxl); all of them mean the same thing to the person who chose it.
    """
    try:
        workbook = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
        return [[_value(cell, percent_as_points) for cell in row]
                for row in workbook.worksheets[0].iter_rows()]
    except Exception:
        return None


def map_headers(header_row, aliases):
    """``{field: column index}`` for the headers recognised, leftmost column winning."""
    mapping = {}
    for index, value in enumerate(header_row):
        label = clean(value).casefold()
        for field, names in aliases.items():
            if field not in mapping and label in names:
                mapping[field] = index
    return mapping


def cell(row, mapping, field):
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def workbook_bytes(headers, rows, widths=None):
    """A small .xlsx held in memory, for the downloadable sample files."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for letter, width in (widths or {}).items():
        sheet.column_dimensions[letter].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
